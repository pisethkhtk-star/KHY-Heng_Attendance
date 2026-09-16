import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from datetime import date, time

def build_attendance_workbook(filepath):
    wb = openpyxl.Workbook()
    
    # Setup color palette matching KHY-HENG HR System
    PRIMARY_DARK = "1E293B"      # Dark Slate for main title / headers
    PRIMARY_BLUE = "2563EB"      # Blue 600
    SUB_HEADER_DARK = "334155"   # Slate 700
    ACCENT_AMBER = "B45309"      # Amber for Late
    ZEBRA_BG = "F8FAFC"          # Slate 50
    CARD_BG = "F1F5F9"           # Slate 100
    BORDER_COLOR = "CBD5E1"      # Slate 300
    HEADER_TEXT_COLOR = "FFFFFF"
    
    font_family = "Segoe UI"
    
    title_font = Font(name=font_family, size=15, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=11, bold=True, color="1E293B")
    header_font = Font(name=font_family, size=10, bold=True, color=HEADER_TEXT_COLOR)
    amber_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10, color="0F172A")
    bold_data_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    kpi_label_font = Font(name=font_family, size=9, bold=True, color="475569")
    
    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )
    
    header_fill = PatternFill(start_color=PRIMARY_DARK, end_color=PRIMARY_DARK, fill_type="solid")
    sub_header_fill = PatternFill(start_color=SUB_HEADER_DARK, end_color=SUB_HEADER_DARK, fill_type="solid")
    blue_header_fill = PatternFill(start_color=PRIMARY_BLUE, end_color=PRIMARY_BLUE, fill_type="solid")
    amber_header_fill = PatternFill(start_color=ACCENT_AMBER, end_color=ACCENT_AMBER, fill_type="solid")
    zebra_fill = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
    
    # -------------------------------------------------------------
    # 1. SETTINGS & LOOKUPS SHEET
    # -------------------------------------------------------------
    ws_settings = wb.active
    ws_settings.title = "Settings"
    ws_settings.views.sheetView[0].showGridLines = True
    
    ws_settings.merge_cells("A1:G1")
    ws_settings["A1"] = "KHY-HENG ATTENDANCE SYSTEM - SETTINGS & CONFIGURATION (ការកំណត់ប្រព័ន្ធ)"
    ws_settings["A1"].font = title_font
    ws_settings["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_settings["A1"].fill = header_fill
    ws_settings.row_dimensions[1].height = 40
    
    ws_settings["A3"] = "COMPANY WORK HOURS & SHIFT RULES (ម៉ោងការងារស្តង់ដារក្រុមហ៊ុន)"
    ws_settings["A3"].font = section_font
    
    shift_headers = ["Parameter Name (ប៉ារ៉ាម៉ែត្រ)", "Value (តម្លៃ)", "Unit / Details (ខ្នាត / បរិយាយ)", "System Code"]
    for col_idx, h in enumerate(shift_headers, start=1):
        cell = ws_settings.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_settings.row_dimensions[4].height = 25
    
    config_rows = [
        ("Morning Shift Start (ម៉ោងចូលវេនព្រឹក)", time(8, 0, 0), "HH:MM (Shift 1 In)", "shift_1_start", "TIME"),
        ("Morning Shift End (ម៉ោងចេញវេនព្រឹក)", time(12, 0, 0), "HH:MM (Shift 1 Out)", "shift_1_end", "TIME"),
        ("Afternoon Shift Start (ម៉ោងចូលវេនរសៀល)", time(13, 0, 0), "HH:MM (Shift 2 In)", "shift_2_start", "TIME"),
        ("Afternoon Shift End (ម៉ោងចេញវេនរសៀល)", time(17, 0, 0), "HH:MM (Shift 2 Out)", "shift_2_end", "TIME"),
        ("Late Grace Period (នាទីអនុគ្រោះយឺត)", 15, "Minutes (អនុគ្រោះយឺត ១៥នាទី)", "late_grace_minutes", "INT"),
        ("Standard Daily Work Hours (ម៉ោងការងារប្រចាំថ្ងៃ)", 8.0, "Hours / Day", "standard_daily_hours", "DEC"),
        ("Standard Monthly Working Days (ថ្ងៃការងារប្រចាំខែ)", 26, "Days / Month", "standard_working_days", "INT")
    ]
    
    for r_idx, row_data in enumerate(config_rows, start=5):
        ws_settings.cell(row=r_idx, column=1, value=row_data[0]).font = bold_data_font
        val_cell = ws_settings.cell(row=r_idx, column=2, value=row_data[1])
        val_cell.font = bold_data_font
        val_cell.alignment = Alignment(horizontal="center")
        if row_data[4] == "TIME":
            val_cell.number_format = "hh:mm"
        elif row_data[4] == "DEC":
            val_cell.number_format = "0.0"
        else:
            val_cell.number_format = "0"
            
        ws_settings.cell(row=r_idx, column=3, value=row_data[2]).font = data_font
        ws_settings.cell(row=r_idx, column=4, value=row_data[3]).font = Font(name=font_family, size=9, color="64748B")
        
        for c in range(1, 5):
            ws_settings.cell(row=r_idx, column=c).border = thin_border
            if r_idx % 2 == 1:
                ws_settings.cell(row=r_idx, column=c).fill = zebra_fill
    
    # Departments Table
    ws_settings["F3"] = "DEPARTMENTS (ផ្នែក)"
    ws_settings["F3"].font = section_font
    dept_headers = ["Dept Code", "Department (EN / KH)"]
    for col_idx, h in enumerate(dept_headers, start=6):
        c = ws_settings.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = sub_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        
    depts = [
        ("IT", "Information Technology (បច្ចេកវិទ្យាព័ត៌មាន)"),
        ("HR", "Human Resources (ធនធានមនុស្ស)"),
        ("FIN", "Finance & Accounting (ហិរញ្ញវត្ថុ និងគណនេយ្យ)"),
        ("MKT", "Marketing & Operations (ទីផ្សារ និងប្រតិបត្តិការ)")
    ]
    for r_idx, (code, name) in enumerate(depts, start=5):
        c1 = ws_settings.cell(row=r_idx, column=6, value=code)
        c2 = ws_settings.cell(row=r_idx, column=7, value=name)
        c1.font = bold_data_font
        c1.alignment = Alignment(horizontal="center")
        c2.font = data_font
        c1.border = thin_border
        c2.border = thin_border
        
    # Leave Types Table
    ws_settings["A14"] = "LEAVE TYPES & ALLOWANCES (ប្រភេទច្បាប់ឈប់សម្រាក)"
    ws_settings["A14"].font = section_font
    leave_headers = ["Code", "Leave Type (EN)", "Leave Type (KH)", "Max Days / Year", "Paid / Unpaid"]
    for col_idx, h in enumerate(leave_headers, start=1):
        c = ws_settings.cell(row=15, column=col_idx, value=h)
        c.font = header_font
        c.fill = sub_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws_settings.row_dimensions[15].height = 25
    
    leave_types = [
        ("AL", "Annual Leave", "ច្បាប់សម្រាកប្រចាំឆ្នាំ", 18, "Paid"),
        ("SL", "Sick Leave", "ច្បាប់ឈឺ", 12, "Paid"),
        ("PL", "Personal Leave", "ច្បាប់ផ្ទាល់ខ្លួន", 7, "Paid"),
        ("ML", "Maternity Leave", "ច្បាប់លំហែមាតុភាព", 90, "Paid (50%-100%)"),
        ("UL", "Unpaid Leave", "ច្បាប់គ្មានប្រាក់ឈ្នួល", 30, "Unpaid")
    ]
    for r_idx, (code, en, kh, days, paid) in enumerate(leave_types, start=16):
        ws_settings.cell(row=r_idx, column=1, value=code).alignment = Alignment(horizontal="center")
        ws_settings.cell(row=r_idx, column=2, value=en)
        ws_settings.cell(row=r_idx, column=3, value=kh)
        d_cell = ws_settings.cell(row=r_idx, column=4, value=days)
        d_cell.alignment = Alignment(horizontal="center")
        p_cell = ws_settings.cell(row=r_idx, column=5, value=paid)
        p_cell.alignment = Alignment(horizontal="center")
        for c in range(1, 6):
            ws_settings.cell(row=r_idx, column=c).border = thin_border
            ws_settings.cell(row=r_idx, column=c).font = data_font
            if r_idx % 2 == 1:
                ws_settings.cell(row=r_idx, column=c).fill = zebra_fill

    # -------------------------------------------------------------
    # 2. EMPLOYEES DIRECTORY SHEET
    # -------------------------------------------------------------
    ws_emp = wb.create_sheet(title="Employees")
    ws_emp.views.sheetView[0].showGridLines = True
    
    ws_emp.merge_cells("A1:K1")
    ws_emp["A1"] = "EMPLOYEE MASTER DIRECTORY (បញ្ជីឈ្មោះបុគ្គលិក)"
    ws_emp["A1"].font = title_font
    ws_emp["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_emp["A1"].fill = header_fill
    ws_emp.row_dimensions[1].height = 40
    
    emp_headers = [
        "Staff ID (អត្តលេខ)",
        "Name in English (ឈ្មោះអង់គ្លេស)",
        "Name in Khmer (ឈ្មោះខ្មែរ)",
        "Gender (ភេទ)",
        "Department (ផ្នែក)",
        "Position (តួនាទី)",
        "Branch (សាខា)",
        "Shift 1 In",
        "Shift 1 Out",
        "Shift 2 In",
        "Shift 2 Out"
    ]
    for col_idx, h in enumerate(emp_headers, start=1):
        c = ws_emp.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_emp.row_dimensions[3].height = 28
    
    sample_employees = [
        ("EMP-001", "Khoem Piseth", "ខឹម ពិសិដ្ឋ", "Male", "Information Technology", "IT Manager", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-002", "Keo Sophea", "កែវ សុភា", "Female", "Human Resources", "HR Manager", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-003", "Chan Dara", "ចាន់ ដារ៉ា", "Male", "Information Technology", "Senior Software Engineer", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-004", "Sok Chenda", "សុខ ចិន្តា", "Female", "Finance & Accounting", "Senior Accountant", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-005", "Meng Chhay", "ម៉េង ឆាយ", "Male", "Marketing & Operations", "Marketing Lead", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-006", "Vannak Bopha", "វណ្ណៈ បុប្ផា", "Female", "Human Resources", "HR Specialist", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-007", "Rithy Visal", "រិទ្ធី វិសាល", "Male", "Information Technology", "UI/UX Designer", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-008", "Chhorn Sreyneang", "ឈន ស្រីនាង", "Female", "Finance & Accounting", "Junior Accountant", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-009", "Long Kosal", "ឡុង កុសល", "Male", "Information Technology", "Network Specialist", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0)),
        ("EMP-010", "Tep Nimol", "ទេព និមល", "Female", "Marketing & Operations", "Operations Officer", "Phnom Penh HQ", time(8,0), time(12,0), time(13,0), time(17,0))
    ]
    
    for r_idx, emp in enumerate(sample_employees, start=4):
        for c_idx, val in enumerate(emp, start=1):
            cell = ws_emp.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx == 1:
                cell.font = bold_data_font
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx >= 8:
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "hh:mm"
                
            if r_idx % 2 == 1:
                cell.fill = zebra_fill
        ws_emp.row_dimensions[r_idx].height = 22

    # -------------------------------------------------------------
    # 3. LEAVES TRACKING SHEET
    # -------------------------------------------------------------
    ws_leaves = wb.create_sheet(title="Leaves")
    ws_leaves.views.sheetView[0].showGridLines = True
    
    ws_leaves.merge_cells("A1:I1")
    ws_leaves["A1"] = "LEAVE REQUESTS & APPROVALS (ការសុំច្បាប់ និងការអនុម័ត)"
    ws_leaves["A1"].font = title_font
    ws_leaves["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_leaves["A1"].fill = header_fill
    ws_leaves.row_dimensions[1].height = 40
    
    leave_cols = [
        "Leave ID",
        "Staff ID",
        "Employee Name (EN)",
        "Leave Type",
        "Date (កាលបរិច្ឆេទ)",
        "Duration (រយៈពេល)",
        "Amount Days",
        "Status (ស្ថានភាព)",
        "Reason / Remarks (មូលហេតុ)"
    ]
    for col_idx, h in enumerate(leave_cols, start=1):
        c = ws_leaves.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_leaves.row_dimensions[3].height = 28
    
    sample_leaves = [
        ("LV-2026-001", "EMP-004", '=IFERROR(VLOOKUP(B4, Employees!$A$4:$C$50, 2, FALSE), "-")', "Annual Leave", date(2026, 9, 3), "Full Day", 1.0, "Approved", "Family event"),
        ("LV-2026-002", "EMP-007", '=IFERROR(VLOOKUP(B5, Employees!$A$4:$C$50, 2, FALSE), "-")', "Sick Leave", date(2026, 9, 3), "Full Day", 1.0, "Approved", "Fever and medical checkup"),
        ("LV-2026-003", "EMP-008", '=IFERROR(VLOOKUP(B6, Employees!$A$4:$C$50, 2, FALSE), "-")', "Sick Leave", date(2026, 9, 1), "Full Day", 1.0, "Approved", "Medical leave approved")
    ]
    for r_idx, lv in enumerate(sample_leaves, start=4):
        for c_idx, val in enumerate(lv, start=1):
            cell = ws_leaves.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if c_idx in [1, 2, 4, 6, 8]:
                cell.alignment = Alignment(horizontal="center")
                if c_idx in [1, 2]:
                    cell.font = bold_data_font
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "yyyy-mm-dd"
            elif c_idx == 7:
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "0.0"
            if r_idx % 2 == 1:
                cell.fill = zebra_fill
        ws_leaves.row_dimensions[r_idx].height = 22

    # -------------------------------------------------------------
    # 4. ATTENDANCE_DAILY SHEET (Core Log & Calculation)
    # -------------------------------------------------------------
    ws_daily = wb.create_sheet(title="Attendance_Daily")
    ws_daily.views.sheetView[0].showGridLines = True
    
    ws_daily.merge_cells("A1:W1")
    ws_daily["A1"] = "DAILY ATTENDANCE LOG & AUTOMATIC COMPUTATION (កំណត់ត្រាវត្តមានប្រចាំថ្ងៃ និងគណនារូបមន្តស្វ័យប្រវត្តិ)"
    ws_daily["A1"].font = title_font
    ws_daily["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_daily["A1"].fill = header_fill
    ws_daily.row_dimensions[1].height = 40
    
    ws_daily.merge_cells("A2:W2")
    ws_daily["A2"] = "✓ រូបមន្តគណនាស្វ័យប្រវត្តិតាមច្បាប់ការងារ: ម៉ោងចូលវេនព្រឹក 08:00 (អនុគ្រោះ 15 នាទី), ម៉ោងចេញវេនព្រឹក 12:00, ម៉ោងចូលវេនរសៀល 13:00 (អនុគ្រោះ 15 នាទី), ម៉ោងចេញវេនរសៀល 17:00"
    ws_daily["A2"].font = Font(name=font_family, size=9, bold=True, color="1E3A8A")
    ws_daily["A2"].fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    ws_daily["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws_daily.row_dimensions[2].height = 24
    
    daily_headers = [
        "No.",
        "Date (កាលបរិច្ឆេទ)",
        "Day of Week",
        "Staff ID (អត្តលេខ)",
        "Employee Name (EN)",
        "Employee Name (KH)",
        "Department (ផ្នែក)",
        "Check-in 1 (ព្រឹកចូល)",
        "Check-out 1 (ព្រឹកចេញ)",
        "Check-in 2 (រសៀលចូល)",
        "Check-out 2 (រសៀលចេញ)",
        "Shift 1 Hours (ម៉ោងវេន១)",
        "Shift 2 Hours (ម៉ោងវេន២)",
        "Total Hours (ម៉ោងសរុប)",
        "Late 1 (នាទីយឺតព្រឹក)",
        "Early 1 (នាទីចេញមុនព្រឹក)",
        "Late 2 (នាទីយឺតរសៀល)",
        "Early 2 (នាទីចេញមុនរសៀល)",
        "Total Late (សរុបយឺត)",
        "Total Early (សរុបចេញមុន)",
        "Daily Status (ស្ថានភាព)",
        "Overtime (OT ម៉ោង)",
        "Note / Remarks (កំណត់សម្គាល់)"
    ]
    
    for col_idx, h in enumerate(daily_headers, start=1):
        c = ws_daily.cell(row=3, column=col_idx, value=h)
        c.font = header_font
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_daily.row_dimensions[3].height = 32
    
    demo_records = [
        # Day 1
        (1, "EMP-001", time(7, 52), time(12, 2), time(12, 55), time(17, 10), 0.0, "Normal on time"),
        (1, "EMP-002", time(7, 58), time(12, 0), time(12, 58), time(17, 2), 0.0, "Normal on time"),
        (1, "EMP-003", time(8, 25), time(12, 5), time(12, 57), time(17, 30), 0.5, "Morning traffic late"),
        (1, "EMP-004", time(7, 50), time(11, 40), time(12, 55), time(17, 0), 0.0, "Early leave morning for errand"),
        (1, "EMP-005", time(8, 5), time(12, 0), time(13, 25), time(17, 5), 0.0, "Late afternoon"),
        (1, "EMP-006", time(8, 20), time(12, 0), time(13, 0), time(16, 45), 0.0, "Late morning & early leave afternoon"),
        (1, "EMP-007", time(7, 55), None, time(13, 0), time(17, 0), 0.0, "Forgot morning check-out"),
        (1, "EMP-008", None, None, None, None, 0.0, "Sick leave (has approval in Leaves sheet)"),
        (1, "EMP-009", None, None, None, None, 0.0, "Unexcused Absence"),
        (1, "EMP-010", time(7, 45), time(12, 0), time(12, 50), time(19, 0), 2.0, "2 hrs Overtime approved"),
        
        # Day 2
        (2, "EMP-001", time(7, 50), time(12, 1), time(12, 52), time(17, 5), 0.0, "On time"),
        (2, "EMP-002", time(7, 55), time(12, 0), time(12, 55), time(17, 0), 0.0, "On time"),
        (2, "EMP-003", time(7, 48), time(12, 0), time(12, 50), time(17, 15), 0.0, "On time"),
        (2, "EMP-004", time(8, 20), time(12, 0), time(13, 0), time(17, 0), 0.0, "Late 5 mins after grace"),
        (2, "EMP-005", time(8, 0), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        (2, "EMP-006", time(8, 0), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        (2, "EMP-007", time(8, 0), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        (2, "EMP-008", time(8, 0), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        (2, "EMP-009", time(8, 10), time(12, 0), time(13, 0), time(17, 0), 0.0, "Within 15m grace - On time"),
        (2, "EMP-010", time(7, 50), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        
        # Day 3
        (3, "EMP-001", time(7, 55), time(12, 5), time(12, 58), time(17, 10), 0.0, "On time"),
        (3, "EMP-002", time(8, 30), time(12, 0), time(13, 0), time(17, 0), 0.0, "Late morning"),
        (3, "EMP-003", time(7, 50), time(12, 0), time(13, 0), time(17, 0), 0.0, "On time"),
        (3, "EMP-004", None, None, None, None, 0.0, "Approved Leave in Leaves sheet"),
        (3, "EMP-005", time(7, 55), time(12, 0), time(12, 55), time(16, 30), 0.0, "Early departure afternoon"),
        (3, "EMP-006", time(7, 58), time(12, 0), time(12, 58), time(17, 0), 0.0, "On time"),
        (3, "EMP-007", None, None, None, None, 0.0, "Approved Sick Leave in Leaves sheet"),
        (3, "EMP-008", time(7, 50), time(12, 0), time(12, 50), time(17, 0), 0.0, "On time"),
        (3, "EMP-009", time(7, 45), time(12, 0), time(12, 45), time(17, 0), 0.0, "On time"),
        (3, "EMP-010", time(7, 55), time(12, 0), time(12, 55), time(18, 30), 1.5, "1.5 hrs Overtime")
    ]
    
    start_r = 4
    for idx, rec in enumerate(demo_records):
        curr_r = start_r + idx
        day_offset, s_id, c1, o1, c2, o2, ot, remark = rec
        rec_date = date(2026, 9, day_offset)
        
        ws_daily.cell(row=curr_r, column=1, value=idx + 1).alignment = Alignment(horizontal="center")
        
        d_cell = ws_daily.cell(row=curr_r, column=2, value=rec_date)
        d_cell.number_format = "yyyy-mm-dd"
        d_cell.alignment = Alignment(horizontal="center")
        
        dow_cell = ws_daily.cell(row=curr_r, column=3, value=f'=TEXT(B{curr_r},"dddd")')
        dow_cell.alignment = Alignment(horizontal="center")
        
        s_cell = ws_daily.cell(row=curr_r, column=4, value=s_id)
        s_cell.font = bold_data_font
        s_cell.alignment = Alignment(horizontal="center")
        
        name_en_cell = ws_daily.cell(row=curr_r, column=5, value=f'=IFERROR(VLOOKUP(D{curr_r}, Employees!$A$4:$F$50, 2, FALSE), "-")')
        name_en_cell.font = bold_data_font
        
        name_kh_cell = ws_daily.cell(row=curr_r, column=6, value=f'=IFERROR(VLOOKUP(D{curr_r}, Employees!$A$4:$F$50, 3, FALSE), "-")')
        name_kh_cell.font = data_font
        
        dept_cell = ws_daily.cell(row=curr_r, column=7, value=f'=IFERROR(VLOOKUP(D{curr_r}, Employees!$A$4:$F$50, 5, FALSE), "-")')
        dept_cell.font = data_font
        
        c1_cell = ws_daily.cell(row=curr_r, column=8, value=c1 if c1 else "")
        c1_cell.alignment = Alignment(horizontal="center")
        if c1: c1_cell.number_format = "hh:mm"
        
        o1_cell = ws_daily.cell(row=curr_r, column=9, value=o1 if o1 else "")
        o1_cell.alignment = Alignment(horizontal="center")
        if o1: o1_cell.number_format = "hh:mm"
        
        c2_cell = ws_daily.cell(row=curr_r, column=10, value=c2 if c2 else "")
        c2_cell.alignment = Alignment(horizontal="center")
        if c2: c2_cell.number_format = "hh:mm"
        
        o2_cell = ws_daily.cell(row=curr_r, column=11, value=o2 if o2 else "")
        o2_cell.alignment = Alignment(horizontal="center")
        if o2: o2_cell.number_format = "hh:mm"
        
        s1_hrs = ws_daily.cell(row=curr_r, column=12, value=f'=IF(AND(H{curr_r}<>"", I{curr_r}<>""), ROUND(MAX(0, (I{curr_r}-H{curr_r})*24), 2), 0)')
        s1_hrs.number_format = "0.00"
        s1_hrs.alignment = Alignment(horizontal="center")
        
        s2_hrs = ws_daily.cell(row=curr_r, column=13, value=f'=IF(AND(J{curr_r}<>"", K{curr_r}<>""), ROUND(MAX(0, (K{curr_r}-J{curr_r})*24), 2), 0)')
        s2_hrs.number_format = "0.00"
        s2_hrs.alignment = Alignment(horizontal="center")
        
        tot_hrs = ws_daily.cell(row=curr_r, column=14, value=f'=L{curr_r}+M{curr_r}')
        tot_hrs.number_format = "0.00"
        tot_hrs.font = bold_data_font
        tot_hrs.alignment = Alignment(horizontal="center")
        
        # Late 1: Shift 1 Start Settings!$B$5 (08:00), Grace Settings!$B$9 (15 mins)
        late1_cell = ws_daily.cell(
            row=curr_r, column=15,
            value=f'=IF(H{curr_r}="","", IF(H{curr_r} > (Settings!$B$5 + Settings!$B$9/1440), ROUND((H{curr_r} - Settings!$B$5)*1440, 0), 0))'
        )
        late1_cell.number_format = "0"
        late1_cell.alignment = Alignment(horizontal="center")
        
        # Early Leave 1: Shift 1 End Settings!$B$6 (12:00)
        early1_cell = ws_daily.cell(
            row=curr_r, column=16,
            value=f'=IF(I{curr_r}="","", IF(I{curr_r} < Settings!$B$6, ROUND((Settings!$B$6 - I{curr_r})*1440, 0), 0))'
        )
        early1_cell.number_format = "0"
        early1_cell.alignment = Alignment(horizontal="center")
        
        # Late 2: Shift 2 Start Settings!$B$7 (13:00), Grace Settings!$B$9 (15 mins)
        late2_cell = ws_daily.cell(
            row=curr_r, column=17,
            value=f'=IF(J{curr_r}="","", IF(J{curr_r} > (Settings!$B$7 + Settings!$B$9/1440), ROUND((J{curr_r} - Settings!$B$7)*1440, 0), 0))'
        )
        late2_cell.number_format = "0"
        late2_cell.alignment = Alignment(horizontal="center")
        
        # Early Leave 2: Shift 2 End Settings!$B$8 (17:00)
        early2_cell = ws_daily.cell(
            row=curr_r, column=18,
            value=f'=IF(K{curr_r}="","", IF(K{curr_r} < Settings!$B$8, ROUND((Settings!$B$8 - K{curr_r})*1440, 0), 0))'
        )
        early2_cell.number_format = "0"
        early2_cell.alignment = Alignment(horizontal="center")
        
        tot_late = ws_daily.cell(row=curr_r, column=19, value=f'=O{curr_r}+Q{curr_r}')
        tot_late.number_format = "0"
        tot_late.alignment = Alignment(horizontal="center")
        tot_late.font = bold_data_font
        
        tot_early = ws_daily.cell(row=curr_r, column=20, value=f'=P{curr_r}+R{curr_r}')
        tot_early.number_format = "0"
        tot_early.alignment = Alignment(horizontal="center")
        tot_early.font = bold_data_font
        
        # Daily Status formula
        status_formula = (
            f'=IF(COUNTIFS(Leaves!$B:$B, D{curr_r}, Leaves!$E:$E, B{curr_r}, Leaves!$H:$H, "Approved")>0, "Leave",'
            f'IF(AND(H{curr_r}="", I{curr_r}="", J{curr_r}="", K{curr_r}=""), "Absent",'
            f'IF(OR(AND(H{curr_r}<>"", I{curr_r}=""), AND(H{curr_r}="", I{curr_r}<>""), AND(J{curr_r}<>"", K{curr_r}=""), AND(J{curr_r}="", K{curr_r}<>"")), "Incomplete Scan",'
            f'IF(AND(S{curr_r}>0, T{curr_r}>0), "Late & Early Leave",'
            f'IF(S{curr_r}>0, "Late",'
            f'IF(T{curr_r}>0, "Early Leave", "On Time"))))))'
        )
        status_cell = ws_daily.cell(row=curr_r, column=21, value=status_formula)
        status_cell.font = bold_data_font
        status_cell.alignment = Alignment(horizontal="center")
        
        ot_cell = ws_daily.cell(row=curr_r, column=22, value=ot)
        ot_cell.number_format = "0.0"
        ot_cell.alignment = Alignment(horizontal="center")
        
        rem_cell = ws_daily.cell(row=curr_r, column=23, value=remark)
        rem_cell.font = data_font
        
        for c in range(1, 24):
            ws_daily.cell(row=curr_r, column=c).border = thin_border
            if curr_r % 2 == 1:
                ws_daily.cell(row=curr_r, column=c).fill = zebra_fill
        ws_daily.row_dimensions[curr_r].height = 22
        
    last_daily_row = start_r + len(demo_records) - 1

    # -------------------------------------------------------------
    # 5. DEDICATED LATE RECORDS LOG SHEET (តារាងកំណត់ត្រាមកយឺតលម្អិត)
    # -------------------------------------------------------------
    ws_late_log = wb.create_sheet(title="Late_Attendance_Log")
    ws_late_log.views.sheetView[0].showGridLines = True
    
    ws_late_log.merge_cells("A1:L1")
    ws_late_log["A1"] = "DETAILED LATE ATTENDANCE AUDIT LOG (កំណត់ត្រាបុគ្គលិកមកយឺតលម្អិត)"
    ws_late_log["A1"].font = title_font
    ws_late_log["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_late_log["A1"].fill = amber_header_fill
    ws_late_log.row_dimensions[1].height = 40
    
    ws_late_log.merge_cells("A2:L2")
    ws_late_log["A2"] = "កំណត់ត្រាជាក់លាក់នៃបុគ្គលិកដែលបានស្កេនចូលយឺតលើសពីម៉ោងកំណត់ និងលើសពី Grace Period 15 នាទី"
    ws_late_log["A2"].font = Font(name=font_family, size=9, italic=True, color="78350F")
    ws_late_log["A2"].fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    ws_late_log["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws_late_log.row_dimensions[2].height = 24
    
    late_log_headers = [
        "No.",
        "Date (កាលបរិច្ឆេទ)",
        "Staff ID",
        "Employee Name (EN)",
        "Employee Name (KH)",
        "Department",
        "Check-in 1",
        "Shift 1 Late (Mins)",
        "Check-in 2",
        "Shift 2 Late (Mins)",
        "Total Late (នាទីសរុប)",
        "Late Analysis Note (បរិយាយភាពយឺត)"
    ]
    for col_idx, h in enumerate(late_log_headers, start=1):
        c = ws_late_log.cell(row=3, column=col_idx, value=h)
        c.font = amber_header_font
        c.fill = amber_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_late_log.row_dimensions[3].height = 28
    
    late_demo_entries = [
        (1, "EMP-003", time(8, 25), 25, None, 0, 25, "Shift 1 Late: 08:25 AM (10 mins after 15m grace)"),
        (1, "EMP-005", None, 0, time(13, 25), 25, 25, "Shift 2 Late: 01:25 PM (10 mins after 15m grace)"),
        (1, "EMP-006", time(8, 20), 20, None, 0, 20, "Shift 1 Late: 08:20 AM (5 mins after 15m grace)"),
        (2, "EMP-004", time(8, 20), 20, None, 0, 20, "Shift 1 Late: 08:20 AM (5 mins after 15m grace)"),
        (3, "EMP-002", time(8, 30), 30, None, 0, 30, "Shift 1 Late: 08:30 AM (15 mins after 15m grace)")
    ]
    for idx, l_rec in enumerate(late_demo_entries):
        lr = 4 + idx
        day_o, s_id, c1, l1, c2, l2, tot_l, note = l_rec
        ws_late_log.cell(row=lr, column=1, value=idx + 1).alignment = Alignment(horizontal="center")
        d_c = ws_late_log.cell(row=lr, column=2, value=date(2026, 9, day_o))
        d_c.number_format = "yyyy-mm-dd"
        d_c.alignment = Alignment(horizontal="center")
        s_c = ws_late_log.cell(row=lr, column=3, value=s_id)
        s_c.font = bold_data_font
        s_c.alignment = Alignment(horizontal="center")
        ws_late_log.cell(row=lr, column=4, value=f'=IFERROR(VLOOKUP(C{lr}, Employees!$A$4:$F$50, 2, FALSE), "-")').font = bold_data_font
        ws_late_log.cell(row=lr, column=5, value=f'=IFERROR(VLOOKUP(C{lr}, Employees!$A$4:$F$50, 3, FALSE), "-")').font = data_font
        ws_late_log.cell(row=lr, column=6, value=f'=IFERROR(VLOOKUP(C{lr}, Employees!$A$4:$F$50, 5, FALSE), "-")').font = data_font
        
        c1_c = ws_late_log.cell(row=lr, column=7, value=c1 if c1 else "-")
        c1_c.alignment = Alignment(horizontal="center")
        if c1: c1_c.number_format = "hh:mm"
        
        l1_c = ws_late_log.cell(row=lr, column=8, value=l1)
        l1_c.alignment = Alignment(horizontal="center")
        l1_c.font = bold_data_font
        
        c2_c = ws_late_log.cell(row=lr, column=9, value=c2 if c2 else "-")
        c2_c.alignment = Alignment(horizontal="center")
        if c2: c2_c.number_format = "hh:mm"
        
        l2_c = ws_late_log.cell(row=lr, column=10, value=l2)
        l2_c.alignment = Alignment(horizontal="center")
        l2_c.font = bold_data_font
        
        tot_c = ws_late_log.cell(row=lr, column=11, value=tot_l)
        tot_c.alignment = Alignment(horizontal="center")
        tot_c.font = Font(name=font_family, size=10, bold=True, color="B45309")
        
        ws_late_log.cell(row=lr, column=12, value=note).font = data_font
        
        for c in range(1, 13):
            ws_late_log.cell(row=lr, column=c).border = thin_border
            if lr % 2 == 1:
                ws_late_log.cell(row=lr, column=c).fill = zebra_fill
        ws_late_log.row_dimensions[lr].height = 22

    # -------------------------------------------------------------
    # 6. ATTENDANCE_REPORT SHEET (With Charts & Table Late)
    # -------------------------------------------------------------
    ws_rep = wb.create_sheet(title="Attendance_Report", index=0)
    ws_rep.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws_rep.merge_cells("A1:P1")
    ws_rep["A1"] = "KHY-HENG ATTENDANCE SYSTEM - ATTENDANCE REPORT & ANALYTICS"
    ws_rep["A1"].font = title_font
    ws_rep["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_rep["A1"].fill = header_fill
    ws_rep.row_dimensions[1].height = 42
    
    ws_rep.merge_cells("A2:P2")
    ws_rep["A2"] = "របាយការណ៍វត្តមាន ស្ថិតិក្រាហ្វិក និងតារាងវិភាគភាពយឺតយ៉ាវ (Attendance Analytics, Visual Charts & Late Tracking Table)"
    ws_rep["A2"].font = Font(name=font_family, size=11, italic=True, color="475569")
    ws_rep["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_rep.row_dimensions[2].height = 22
    
    # 5 KPI Metric Cards (Row 4 to 6)
    kpi_configs = [
        ("A", "C", "TOTAL EMPLOYEES (បុគ្គលិកសរុប)", "=COUNTA(Employees!A4:A50)", "2563EB"),
        ("D", "F", "LOGGED ATTENDANCE (វត្តមានកត់ត្រា)", f"=COUNTA(Attendance_Daily!A4:A{last_daily_row})", "0F172A"),
        ("G", "I", "ON TIME (ទាន់ពេល)", f'=COUNTIF(Attendance_Daily!U4:U{last_daily_row}, "On Time")', "059669"),
        ("J", "L", "LATE ARRIVALS (ការមកយឺត)", f'=COUNTIF(Attendance_Daily!S4:S{last_daily_row}, ">0")', "D97706"),
        ("M", "P", "EARLY / ABSENT (ចេញមុន & អវត្តមាន)", f'=COUNTIF(Attendance_Daily!U4:U{last_daily_row}, "*Absent*") + COUNTIF(Attendance_Daily!T4:T{last_daily_row}, ">0")', "DC2626")
    ]
    
    for c_start, c_end, label, formula, color_hex in kpi_configs:
        ws_rep.merge_cells(f"{c_start}4:{c_end}4")
        lbl_cell = ws_rep[f"{c_start}4"]
        lbl_cell.value = label
        lbl_cell.font = kpi_label_font
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        ws_rep.merge_cells(f"{c_start}5:{c_end}6")
        val_c = ws_rep[f"{c_start}5"]
        val_c.value = formula
        val_c.font = Font(name=font_family, size=20, bold=True, color=color_hex)
        val_c.alignment = Alignment(horizontal="center", vertical="center")
        
        start_idx = openpyxl.utils.column_index_from_string(c_start)
        end_idx = openpyxl.utils.column_index_from_string(c_end)
        for r in range(4, 7):
            for c_i in range(start_idx, end_idx + 1):
                cell = ws_rep.cell(row=r, column=c_i)
                cell.fill = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")
                cell.border = thin_border
                
    ws_rep.row_dimensions[4].height = 18
    ws_rep.row_dimensions[5].height = 18
    ws_rep.row_dimensions[6].height = 18
    
    # -------------------------------------------------------------
    # CHART DATA TABLES (Columns N to O, Rows 8 to 13)
    # -------------------------------------------------------------
    ws_rep["N8"] = "Status Category"
    ws_rep["O8"] = "Record Count"
    ws_rep["N8"].font = header_font
    ws_rep["O8"].font = header_font
    ws_rep["N8"].fill = sub_header_fill
    ws_rep["O8"].fill = sub_header_fill
    ws_rep["N8"].border = thin_border
    ws_rep["O8"].border = thin_border
    
    chart_stat_rows = [
        ("On Time (ទាន់ពេល)", f'=COUNTIF(Attendance_Daily!$U$4:$U${last_daily_row}, "On Time")'),
        ("Late (មកយឺត)", f'=COUNTIF(Attendance_Daily!$S$4:$S${last_daily_row}, ">0")'),
        ("Early Leave (ចេញមុន)", f'=COUNTIF(Attendance_Daily!$T$4:$T${last_daily_row}, ">0")'),
        ("Leave (ច្បាប់)", f'=COUNTIF(Attendance_Daily!$U$4:$U${last_daily_row}, "*Leave*")'),
        ("Absent (អវត្តមាន)", f'=COUNTIF(Attendance_Daily!$U$4:$U${last_daily_row}, "Absent")')
    ]
    for idx, (stat_name, stat_form) in enumerate(chart_stat_rows, start=9):
        c1 = ws_rep.cell(row=idx, column=14, value=stat_name)
        c2 = ws_rep.cell(row=idx, column=15, value=stat_form)
        c1.font = data_font
        c2.font = bold_data_font
        c2.alignment = Alignment(horizontal="center")
        c1.border = thin_border
        c2.border = thin_border
        if idx % 2 == 1:
            c1.fill = zebra_fill
            c2.fill = zebra_fill
            
    # Add Doughnut Chart for Attendance Breakdown
    chart_doughnut = DoughnutChart()
    chart_doughnut.title = "Attendance Status Distribution (ភាគរយវត្តមាន)"
    chart_doughnut.style = 10
    chart_doughnut.width = 15.0
    chart_doughnut.height = 9.2
    
    data_donut = Reference(ws_rep, min_col=15, min_row=8, max_row=13)
    cats_donut = Reference(ws_rep, min_col=14, min_row=9, max_row=13)
    chart_doughnut.add_data(data_donut, titles_from_data=True)
    chart_doughnut.set_categories(cats_donut)
    chart_doughnut.dataLabels = DataLabelList()
    chart_doughnut.dataLabels.showVal = True
    chart_doughnut.dataLabels.showPercent = True
    ws_rep.add_chart(chart_doughnut, "A8")

    # -------------------------------------------------------------
    # TABLE LATE: LATE ATTENDANCE & PUNCTUALITY ANALYSIS
    # -------------------------------------------------------------
    ws_rep["A24"] = "1. LATE ATTENDANCE & PUNCTUALITY ANALYSIS TABLE (តារាងតាមដាន និងវិភាគភាពយឺតយ៉ាវរបស់បុគ្គលិក)"
    ws_rep["A24"].font = Font(name=font_family, size=12, bold=True, color="92400E")
    ws_rep.row_dimensions[24].height = 26
    
    late_table_headers = [
        "Staff ID (អត្តលេខ)",
        "Employee Name (EN)",
        "Employee Name (KH)",
        "Department (ផ្នែក)",
        "Scheduled Days",
        "Late Count (ដងយឺត)",
        "Morning Late (នាទីព្រឹក)",
        "Afternoon Late (នាទីរសៀល)",
        "Total Late Mins (សរុបនាទីយឺត)",
        "Avg Late (មធ្យម/ដង)",
        "Exceeded Grace (>15m)",
        "Late Status / Action (កម្រិតចំណាត់ការ)"
    ]
    
    for col_idx, h in enumerate(late_table_headers, start=1):
        c = ws_rep.cell(row=25, column=col_idx, value=h)
        c.font = amber_header_font
        c.fill = amber_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_rep.row_dimensions[25].height = 30
    
    late_start_r = 26
    for idx in range(len(sample_employees)):
        lr = late_start_r + idx
        emp_r = 4 + idx
        
        # Staff ID
        s_c = ws_rep.cell(row=lr, column=1, value=f'=Employees!A{emp_r}')
        s_c.font = bold_data_font
        s_c.alignment = Alignment(horizontal="center")
        
        # Name EN
        ws_rep.cell(row=lr, column=2, value=f'=Employees!B{emp_r}').font = bold_data_font
        # Name KH
        ws_rep.cell(row=lr, column=3, value=f'=Employees!C{emp_r}').font = data_font
        # Department
        ws_rep.cell(row=lr, column=4, value=f'=Employees!E{emp_r}').font = data_font
        
        # Scheduled Days
        ws_rep.cell(row=lr, column=5, value="=Settings!$B$11").alignment = Alignment(horizontal="center")
        
        # Late Count = COUNTIFS(StaffId, TotalLate > 0)
        lc_c = ws_rep.cell(row=lr, column=6, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{lr}, Attendance_Daily!$S$4:$S${last_daily_row}, ">0")')
        lc_c.alignment = Alignment(horizontal="center")
        lc_c.font = Font(name=font_family, size=10, bold=True, color="B45309")
        lc_c.number_format = "0"
        
        # Morning Late Mins = SUMIFS(Late 1, StaffId)
        m_late = ws_rep.cell(row=lr, column=7, value=f'=SUMIFS(Attendance_Daily!$O$4:$O${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{lr})')
        m_late.alignment = Alignment(horizontal="center")
        m_late.number_format = "0"
        
        # Afternoon Late Mins = SUMIFS(Late 2, StaffId)
        a_late = ws_rep.cell(row=lr, column=8, value=f'=SUMIFS(Attendance_Daily!$Q$4:$Q${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{lr})')
        a_late.alignment = Alignment(horizontal="center")
        a_late.number_format = "0"
        
        # Total Late Minutes = SUMIFS(TotalLate, StaffId)
        tot_late_c = ws_rep.cell(row=lr, column=9, value=f'=SUMIFS(Attendance_Daily!$S$4:$S${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{lr})')
        tot_late_c.alignment = Alignment(horizontal="center")
        tot_late_c.font = Font(name=font_family, size=10, bold=True, color="92400E")
        tot_late_c.number_format = "0"
        
        # Avg Late Mins/Occurrence = IF(LateCount>0, TotalLate/LateCount, 0)
        avg_late = ws_rep.cell(row=lr, column=10, value=f'=IF(F{lr}>0, ROUND(I{lr}/F{lr}, 1), 0)')
        avg_late.alignment = Alignment(horizontal="center")
        avg_late.number_format = "0.0"
        
        # Exceeded Grace (>15m)
        exc_c = ws_rep.cell(row=lr, column=11, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{lr}, Attendance_Daily!$S$4:$S${last_daily_row}, ">15")')
        exc_c.alignment = Alignment(horizontal="center")
        exc_c.number_format = "0"
        
        # Late Status / Warning Level
        status_action_formula = (
            f'=IF(I{lr}=0, "Punctual / ទៀងទាត់",'
            f'IF(F{lr}<=1, "Acceptable / អនុគ្រោះ",'
            f'IF(F{lr}<=2, "Verbal Warning / ព្រមានផ្ទាល់មាត់",'
            f'"Formal Warning / ព្រមានជាលាយលក្ខណ៍អក្សរ")))'
        )
        act_c = ws_rep.cell(row=lr, column=12, value=status_action_formula)
        act_c.font = bold_data_font
        act_c.alignment = Alignment(horizontal="center")
        
        for c in range(1, 13):
            ws_rep.cell(row=lr, column=c).border = thin_border
            if lr % 2 == 1:
                ws_rep.cell(row=lr, column=c).fill = zebra_fill
        ws_rep.row_dimensions[lr].height = 22
        
    late_last_r = late_start_r + len(sample_employees) - 1
    
    # Total Row for Late Table
    late_tot_r = late_last_r + 1
    ws_rep.cell(row=late_tot_r, column=1, value="TOTAL LATE (សរុប)").alignment = Alignment(horizontal="center")
    ws_rep.cell(row=late_tot_r, column=1).font = amber_header_font
    for c in range(1, 13):
        cell = ws_rep.cell(row=late_tot_r, column=c)
        cell.fill = amber_header_fill
        cell.border = thin_border
        cell.font = amber_header_font
        if c in [6, 7, 8, 9, 11]:
            col_ltr = get_column_letter(c)
            cell.value = f"=SUM({col_ltr}{late_start_r}:{col_ltr}{late_last_r})"
            cell.alignment = Alignment(horizontal="center")
            cell.number_format = "0"
    ws_rep.row_dimensions[late_tot_r].height = 24
    
    # Add Bar Chart for Late Minutes by Employee
    chart_bar = BarChart()
    chart_bar.type = "col"
    chart_bar.style = 10
    chart_bar.title = "Total Late Minutes by Employee (នាទីមកយឺតតាមបុគ្គលិក)"
    chart_bar.y_axis.title = "Late Minutes (នាទី)"
    chart_bar.x_axis.title = "Employee (បុគ្គលិក)"
    chart_bar.width = 16.0
    chart_bar.height = 9.2
    
    # Data is column I (Total Late Mins) from Late Table, categories is column B (Employee Name EN)
    data_bar = Reference(ws_rep, min_col=9, min_row=25, max_row=late_last_r)
    cats_bar = Reference(ws_rep, min_col=2, min_row=26, max_row=late_last_r)
    chart_bar.add_data(data_bar, titles_from_data=True)
    chart_bar.set_categories(cats_bar)
    chart_bar.legend = None
    ws_rep.add_chart(chart_bar, "G8")

    # -------------------------------------------------------------
    # OVERALL ATTENDANCE SUMMARY TABLE
    # -------------------------------------------------------------
    sum_section_r = late_tot_r + 2
    ws_rep.cell(row=sum_section_r, column=1, value="2. OVERALL ATTENDANCE SUMMARY TABLE (តារាងសង្ខេបវត្តមានទូទៅប្រចាំខែ)").font = section_font
    ws_rep.row_dimensions[sum_section_r].height = 26
    
    sum_header_r = sum_section_r + 1
    summary_headers = [
        "Staff ID (អត្តលេខ)",
        "Employee Name (EN)",
        "Employee Name (KH)",
        "Department (ផ្នែក)",
        "Scheduled Days",
        "Present Days (វត្តមាន)",
        "On Time (ទាន់ពេល)",
        "Late Count (ដងយឺត)",
        "Total Late Mins (នាទីយឺត)",
        "Early Count (ដងចេញមុន)",
        "Total Early Mins (នាទីចេញមុន)",
        "Leave Days (ច្បាប់)",
        "Absent Days (អវត្តមាន)",
        "Total Work Hours (ម៉ោងការងារ)",
        "Overtime Hours (OT)"
    ]
    
    for col_idx, h in enumerate(summary_headers, start=1):
        c = ws_rep.cell(row=sum_header_r, column=col_idx, value=h)
        c.font = header_font
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws_rep.row_dimensions[sum_header_r].height = 30
    
    sum_data_start = sum_header_r + 1
    for idx in range(len(sample_employees)):
        sum_r = sum_data_start + idx
        emp_r = 4 + idx
        
        # Col 1: Staff ID
        s_cell = ws_rep.cell(row=sum_r, column=1, value=f'=Employees!A{emp_r}')
        s_cell.font = bold_data_font
        s_cell.alignment = Alignment(horizontal="center")
        
        # Col 2: Name EN
        ws_rep.cell(row=sum_r, column=2, value=f'=Employees!B{emp_r}').font = bold_data_font
        # Col 3: Name KH
        ws_rep.cell(row=sum_r, column=3, value=f'=Employees!C{emp_r}').font = data_font
        # Col 4: Department
        ws_rep.cell(row=sum_r, column=4, value=f'=Employees!E{emp_r}').font = data_font
        
        # Col 5: Scheduled Days
        ws_rep.cell(row=sum_r, column=5, value="=Settings!$B$11").alignment = Alignment(horizontal="center")
        
        # Col 6: Present Days
        p_cell = ws_rep.cell(row=sum_r, column=6, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$U$4:$U${last_daily_row}, "<>Absent", Attendance_Daily!$U$4:$U${last_daily_row}, "<>Leave")')
        p_cell.alignment = Alignment(horizontal="center")
        p_cell.font = bold_data_font
        p_cell.number_format = "0"
        
        # Col 7: On Time Days
        ot_cell = ws_rep.cell(row=sum_r, column=7, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$U$4:$U${last_daily_row}, "On Time")')
        ot_cell.alignment = Alignment(horizontal="center")
        ot_cell.font = Font(name=font_family, size=10, bold=True, color="059669")
        ot_cell.number_format = "0"
        
        # Col 8: Late Count
        lc_cell = ws_rep.cell(row=sum_r, column=8, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$S$4:$S${last_daily_row}, ">0")')
        lc_cell.alignment = Alignment(horizontal="center")
        lc_cell.font = Font(name=font_family, size=10, bold=True, color="D97706")
        lc_cell.number_format = "0"
        
        # Col 9: Total Late Mins
        tl_cell = ws_rep.cell(row=sum_r, column=9, value=f'=SUMIFS(Attendance_Daily!$S$4:$S${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r})')
        tl_cell.alignment = Alignment(horizontal="center")
        tl_cell.number_format = "0"
        
        # Col 10: Early Count
        ec_cell = ws_rep.cell(row=sum_r, column=10, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$T$4:$T${last_daily_row}, ">0")')
        ec_cell.alignment = Alignment(horizontal="center")
        ec_cell.font = Font(name=font_family, size=10, bold=True, color="EA580C")
        ec_cell.number_format = "0"
        
        # Col 11: Total Early Mins
        te_cell = ws_rep.cell(row=sum_r, column=11, value=f'=SUMIFS(Attendance_Daily!$T$4:$T${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r})')
        te_cell.alignment = Alignment(horizontal="center")
        te_cell.number_format = "0"
        
        # Col 12: Leave Days
        lv_cell = ws_rep.cell(row=sum_r, column=12, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$U$4:$U${last_daily_row}, "*Leave*")')
        lv_cell.alignment = Alignment(horizontal="center")
        lv_cell.font = Font(name=font_family, size=10, color="2563EB")
        lv_cell.number_format = "0"
        
        # Col 13: Absent Days
        ab_cell = ws_rep.cell(row=sum_r, column=13, value=f'=COUNTIFS(Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r}, Attendance_Daily!$U$4:$U${last_daily_row}, "Absent")')
        ab_cell.alignment = Alignment(horizontal="center")
        ab_cell.font = Font(name=font_family, size=10, bold=True, color="DC2626")
        ab_cell.number_format = "0"
        
        # Col 14: Total Work Hours
        th_cell = ws_rep.cell(row=sum_r, column=14, value=f'=SUMIFS(Attendance_Daily!$N$4:$N${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r})')
        th_cell.alignment = Alignment(horizontal="center")
        th_cell.font = bold_data_font
        th_cell.number_format = "0.00"
        
        # Col 15: Overtime Hours
        ot_tot_cell = ws_rep.cell(row=sum_r, column=15, value=f'=SUMIFS(Attendance_Daily!$V$4:$V${last_daily_row}, Attendance_Daily!$D$4:$D${last_daily_row}, A{sum_r})')
        ot_tot_cell.alignment = Alignment(horizontal="center")
        ot_tot_cell.font = bold_data_font
        ot_tot_cell.number_format = "0.0"
        
        for c in range(1, 16):
            ws_rep.cell(row=sum_r, column=c).border = thin_border
            if sum_r % 2 == 1:
                ws_rep.cell(row=sum_r, column=c).fill = zebra_fill
        ws_rep.row_dimensions[sum_r].height = 22
        
    sum_last_r = sum_data_start + len(sample_employees) - 1
    
    # Overall Total Footer Row
    tot_r = sum_last_r + 1
    ws_rep.cell(row=tot_r, column=1, value="OVERALL TOTAL (សរុបរួម)").alignment = Alignment(horizontal="center")
    ws_rep.cell(row=tot_r, column=1).font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    for c in range(1, 16):
        cell = ws_rep.cell(row=tot_r, column=c)
        cell.fill = header_fill
        cell.border = thin_border
        cell.font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
        if c >= 6:
            col_ltr = get_column_letter(c)
            cell.value = f"=SUM({col_ltr}{sum_data_start}:{col_ltr}{sum_last_r})"
            cell.alignment = Alignment(horizontal="center")
            if c in [14]:
                cell.number_format = "0.00"
            elif c in [15]:
                cell.number_format = "0.0"
            else:
                cell.number_format = "0"
    ws_rep.row_dimensions[tot_r].height = 25

    # -------------------------------------------------------------
    # 7. DATA VALIDATION (DROPDOWNS)
    # -------------------------------------------------------------
    dv_staff = DataValidation(type="list", formula1="=Employees!$A$4:$A$50", allow_blank=True)
    ws_daily.add_data_validation(dv_staff)
    dv_staff.add(f"D4:D{last_daily_row + 200}")
    
    dv_leave_dur = DataValidation(type="list", formula1='"Full Day,Morning Shift,Afternoon Shift"', allow_blank=True)
    ws_leaves.add_data_validation(dv_leave_dur)
    dv_leave_dur.add("F4:F100")
    
    dv_leave_status = DataValidation(type="list", formula1='"Approved,Pending,Rejected"', allow_blank=True)
    ws_leaves.add_data_validation(dv_leave_status)
    dv_leave_status.add("H4:H100")
    
    dv_leave_type = DataValidation(type="list", formula1="=Settings!$B$16:$B$20", allow_blank=True)
    ws_leaves.add_data_validation(dv_leave_type)
    dv_leave_type.add("D4:D100")

    # -------------------------------------------------------------
    # 8. CONDITIONAL FORMATTING
    # -------------------------------------------------------------
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    green_font = Font(name=font_family, size=10, bold=True, color="065F46")
    
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    yellow_font = Font(name=font_family, size=10, bold=True, color="92400E")
    
    orange_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    orange_font = Font(name=font_family, size=10, bold=True, color="9A3412")
    
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name=font_family, size=10, bold=True, color="991B1B")
    
    purple_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid")
    purple_font = Font(name=font_family, size=10, bold=True, color="6B21A8")
    
    blue_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    blue_font = Font(name=font_family, size=10, bold=True, color="1E40AF")
    
    status_range = f"U4:U{last_daily_row + 100}"
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"On Time"'], fill=green_fill, font=green_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Late"'], fill=yellow_fill, font=yellow_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Early Leave"'], fill=orange_fill, font=orange_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Late & Early Leave"'], fill=yellow_fill, font=yellow_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Incomplete Scan"'], fill=purple_fill, font=purple_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Absent"'], fill=red_fill, font=red_font))
    ws_daily.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Leave"'], fill=blue_fill, font=blue_font))
    
    ws_daily.conditional_formatting.add(f"S4:S{last_daily_row + 100}", CellIsRule(operator='greaterThan', formula=['0'], fill=yellow_fill, font=yellow_font))
    ws_daily.conditional_formatting.add(f"T4:T{last_daily_row + 100}", CellIsRule(operator='greaterThan', formula=['0'], fill=orange_fill, font=orange_font))
    
    # Conditional formatting on Late Table in Attendance_Report
    ws_rep.conditional_formatting.add(f"F{late_start_r}:F{late_last_r}", CellIsRule(operator='greaterThan', formula=['0'], fill=yellow_fill, font=yellow_font))
    ws_rep.conditional_formatting.add(f"I{late_start_r}:I{late_last_r}", CellIsRule(operator='greaterThan', formula=['0'], fill=orange_fill, font=orange_font))
    ws_rep.conditional_formatting.add(f"L{late_start_r}:L{late_last_r}", CellIsRule(operator='equal', formula=['"Formal Warning / ព្រមានជាលាយលក្ខណ៍អក្សរ"'], fill=red_fill, font=red_font))

    # -------------------------------------------------------------
    # 9. USER GUIDE SHEET (របៀបប្រើប្រាស់)
    # -------------------------------------------------------------
    ws_guide = wb.create_sheet(title="User_Guide")
    ws_guide.views.sheetView[0].showGridLines = True
    
    ws_guide.merge_cells("A1:H1")
    ws_guide["A1"] = "USER INSTRUCTIONS & SYSTEM FORMULAS GUIDE (ការណែនាំអំពីរបៀបប្រើប្រាស់ និងរូបមន្ត)"
    ws_guide["A1"].font = title_font
    ws_guide["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_guide["A1"].fill = header_fill
    ws_guide.row_dimensions[1].height = 40
    
    guide_content = [
        ("១. ទិដ្ឋភាពទូទៅនៃប្រព័ន្ធ (System Overview)", 
         "File Excel នេះត្រូវបានរចនាឡើងយ៉ាងម៉ត់ចត់ដោយផ្អែកលើ logic នៃ KHY-HENG Attendance System ទាំង backend (Spring Boot) និង frontend (React) ដើម្បីគ្រប់គ្រងវត្តមាន ម៉ោងការងារ ភាពយឺតយ៉ាវ ការចេញមុន ច្បាប់សម្រាក និងម៉ោងបន្ថែម (OT) តាមស្តង់ដារជាក់ស្តែង។"),
        
        ("២. ផ្ទាំងរបាយការណ៍ និងក្រាហ្វិក (Attendance_Report Sheet)", 
         "• ផ្នែកខាងលើមាន KPI Cards សំខាន់ៗចំនួន ៥ (Total Staff, Logged Days, On Time, Late Arrivals, Early/Absent)\n• មាន Charts ក្រាហ្វិកគំនូសតាងចំនួន ២:\n   1. Doughnut Chart: បង្ហាញភាគរយវត្តមានតាមប្រភេទ (On Time, Late, Early Leave, Leave, Absent)\n   2. Bar Chart: បង្ហាញការប្រៀបធៀបនាទីមកយឺតសរុបតាមបុគ្គលិកនីមួយៗ (Total Late Minutes by Employee)\n• មាន Table Late (តារាងតាមដាន និងវិភាគភាពយឺតយ៉ាវ): បង្ហាញចំនួនដងយឺត នាទីយឺតវេនព្រឹក នាទីយឺតរសៀល មធ្យមភាគនាទីយឺត និងកម្រិតព្រមាន\n• មាន Overall Attendance Summary Table: បង្ហាញសរុបថ្ងៃវត្តមាន ម៉ោងការងារ និង OT ប្រចាំខែ។"),
        
        ("៣. កំណត់ត្រាមកយឺតលម្អិត (Late_Attendance_Log Sheet)", 
         "• សន្លឹកនេះកត់ត្រាគ្រប់ករណីស្កេនយឺតជាក់ស្តែងទាំងអស់ ដោយបង្ហាញម៉ោងស្កេនចូល នាទីយឺតធៀបនឹង Grace Period ១៥នាទី និងកំណត់ចំណាំលម្អិតដូចក្នុងផ្ទាំង AttendanceLate នៃប្រព័ន្ធ frontend។"),
        
        ("៤. ការកំណត់ម៉ោង និងការអនុគ្រោះ (Settings Sheet)", 
         "• វេនព្រឹក (Morning Shift): 08:00 - 12:00\n• វេនរសៀល (Afternoon Shift): 13:00 - 17:00\n• ការអនុគ្រោះយឺត (Late Grace Period): ១៥ នាទី (ប្រសិនបើចូលមុន 08:15 ឬ 13:15 មិនគិតថាយឺតទេ)\n• អ្នកអាចផ្លាស់ប្តូរម៉ោងស្តង់ដារក្នុងសន្លឹក 'Settings' បានគ្រប់ពេល រូបមន្តទាំងអស់នឹង update ដោយស្វ័យប្រវត្តិ។"),
        
        ("៥. ការកត់ត្រាវត្តមានប្រចាំថ្ងៃ (Attendance_Daily Sheet)", 
         "• បញ្ចូលកាលបរិច្ឆេទ និង Staff ID (មាន dropdown ងាយស្រួលជ្រើសរើស)\n• បញ្ចូលម៉ោង Check-in 1, Check-out 1, Check-in 2, Check-out 2\n• ប្រព័ន្ធនឹងគណនាម៉ោងការងារ នាទីមកយឺត នាទីចេញមុន និងស្ថានភាព (On Time, Late, Early Leave, Incomplete Scan, Absent, Leave) ដោយស្វ័យប្រវត្តិ។")
    ]
    
    g_row = 3
    for title, desc in guide_content:
        ws_guide.merge_cells(start_row=g_row, start_column=1, end_row=g_row, end_column=8)
        c_title = ws_guide.cell(row=g_row, column=1, value=title)
        c_title.font = section_font
        c_title.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws_guide.row_dimensions[g_row].height = 24
        
        g_row += 1
        ws_guide.merge_cells(start_row=g_row, start_column=1, end_row=g_row, end_column=8)
        c_desc = ws_guide.cell(row=g_row, column=1, value=desc)
        c_desc.font = data_font
        c_desc.alignment = Alignment(wrap_text=True, vertical="top")
        
        line_count = desc.count('\n') + 2
        ws_guide.row_dimensions[g_row].height = line_count * 18
        
        g_row += 2

    # -------------------------------------------------------------
    # 10. AUTO-FIT COLUMN WIDTHS
    # -------------------------------------------------------------
    for ws in [ws_rep, ws_late_log, ws_daily, ws_emp, ws_leaves, ws_settings, ws_guide]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row in [1, 2] or (ws.title == "User_Guide" and cell.row > 2):
                    continue
                val_str = str(cell.value or '')
                if val_str.startswith('='):
                    max_len = max(max_len, 12)
                else:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    # Tailored column widths for Attendance_Report
    ws_rep.column_dimensions["A"].width = 13
    ws_rep.column_dimensions["B"].width = 22
    ws_rep.column_dimensions["C"].width = 20
    ws_rep.column_dimensions["D"].width = 25
    ws_rep.column_dimensions["E"].width = 15
    ws_rep.column_dimensions["F"].width = 16
    ws_rep.column_dimensions["G"].width = 16
    ws_rep.column_dimensions["H"].width = 16
    ws_rep.column_dimensions["I"].width = 16
    ws_rep.column_dimensions["J"].width = 15
    ws_rep.column_dimensions["K"].width = 16
    ws_rep.column_dimensions["L"].width = 26
    ws_rep.column_dimensions["M"].width = 15
    ws_rep.column_dimensions["N"].width = 16
    ws_rep.column_dimensions["O"].width = 16
    ws_rep.column_dimensions["P"].width = 16
    
    # Tailored column widths for Late_Attendance_Log
    ws_late_log.column_dimensions["B"].width = 14
    ws_late_log.column_dimensions["C"].width = 13
    ws_late_log.column_dimensions["D"].width = 22
    ws_late_log.column_dimensions["E"].width = 20
    ws_late_log.column_dimensions["F"].width = 24
    ws_late_log.column_dimensions["K"].width = 16
    ws_late_log.column_dimensions["L"].width = 38
    
    ws_daily.column_dimensions["B"].width = 13
    ws_daily.column_dimensions["C"].width = 14
    ws_daily.column_dimensions["D"].width = 13
    ws_daily.column_dimensions["E"].width = 20
    ws_daily.column_dimensions["F"].width = 18
    ws_daily.column_dimensions["G"].width = 24
    ws_daily.column_dimensions["U"].width = 22
    ws_daily.column_dimensions["W"].width = 30
    
    ws_emp.column_dimensions["B"].width = 22
    ws_emp.column_dimensions["C"].width = 20
    ws_emp.column_dimensions["E"].width = 25
    ws_emp.column_dimensions["F"].width = 25
    
    # Save workbook with lock fallback
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    try:
        wb.save(filepath)
        print(f"Successfully generated: {filepath}")
    except PermissionError:
        alt_path = filepath.replace(".xlsx", "_Updated.xlsx")
        wb.save(alt_path)
        print(f"Original file is currently open in Excel! Saved to: {alt_path}")

if __name__ == "__main__":
    out_path = r"d:\project\KHY-HENG_Attendance\KHY_HENG_Attendance_Management_System.xlsx"
    build_attendance_workbook(out_path)
